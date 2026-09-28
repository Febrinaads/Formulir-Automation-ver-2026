"""
pptx_builder.py
Membuat slide "Form Penjelasan Dokumentasi" (A4 landscape) memakai python-pptx.
Dipisah dari app.py supaya bisa dites / dipakai tanpa Streamlit.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import List, Optional

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x5C)
SOFT = RGBColor(0xDD, 0xEB, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x11, 0x11, 0x11)
RED = RGBColor(0xFF, 0x00, 0x00)
YELLOW = RGBColor(0xFF, 0xFF, 0x00)
LINK = RGBColor(0x05, 0x63, 0xC1)
FONT = "Arial"


# ----------------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------------
@dataclass
class ImageItem:
    data: bytes          # isi file gambar (png/jpg/...)
    caption: str = ""    # deskripsi (maks. 20 kata)


@dataclass
class FormData:
    jenis_kegiatan: str
    nama: str
    nik: str
    jabatan: str
    email: str
    hsh: str
    perusahaan: str
    direktorat: str
    tanggal: str
    waktu_mulai: str     # "15:00"
    waktu_selesai: str   # "15:00"
    topik: str
    uraian: str
    images: List[ImageItem] = field(default_factory=list)
    logo_kiri: Optional[bytes] = None
    logo_kanan: Optional[bytes] = None


def count_words(text: str) -> int:
    return len(re.findall(r"\S+", text or ""))


def strip_markup(text: str) -> str:
    return (text or "").replace("**", "")


# ----------------------------------------------------------------------------
# Helper gambar
# ----------------------------------------------------------------------------
def _prepare_image(raw: bytes, max_dim: int = 1600, keep_alpha: bool = False):
    """Perkecil gambar & kembalikan (bytes, width_px, height_px)."""
    im = Image.open(io.BytesIO(raw))
    im.load()
    if max(im.size) > max_dim:
        r = max_dim / max(im.size)
        im = im.resize((round(im.width * r), round(im.height * r)), Image.LANCZOS)
    out = io.BytesIO()
    if keep_alpha:
        im.convert("RGBA").save(out, format="PNG")
    else:
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        im.save(out, format="JPEG", quality=88)
    return out.getvalue(), im.width, im.height


def _fit(w, h, bx, by, bw, bh, halign="center"):
    """Skala (w,h) agar muat di kotak (bx,by,bw,bh); halign: left/center/right."""
    r = min(bw / w, bh / h)
    nw, nh = w * r, h * r
    dx = {"left": 0, "center": (bw - nw) / 2, "right": bw - nw}[halign]
    return bx + dx, by + (bh - nh) / 2, nw, nh


# ----------------------------------------------------------------------------
# Helper shape / teks (semua ukuran dalam inci)
# ----------------------------------------------------------------------------
def _rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    # buang referensi style tema (efek bayangan/garis bawaan) agar tampilan datar
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    return shp


def _textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE, margin=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    return tb, tf


def _run(paragraph, text, size=10, bold=False, color=BLACK, underline=False, link=None):
    r = paragraph.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.underline = underline
    f.color.rgb = color
    if link:
        r.hyperlink.address = link
    return r


def _simple_text(slide, text, x, y, w, h, size=10, bold=False, color=BLACK,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    _, tf = _textbox(slide, x, y, w, h, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, bold=bold, color=color)


def _row(slide, label, value, x, y, label_w, value_w, email=False):
    _simple_text(slide, label, x, y, label_w, 0.26)
    _, tf = _textbox(slide, x + label_w, y, value_w, 0.26)
    p = tf.paragraphs[0]
    _run(p, ": ")
    if email and value:
        _run(p, value, color=LINK, underline=True, link=f"mailto:{value}")
    else:
        _run(p, value)


def _add_markup_runs(paragraph, line: str, size: float):
    """Teks dengan **bagian merah**."""
    for part in re.split(r"(\*\*[^*]+\*\*)", line):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _run(paragraph, part[2:-2], size=size, color=RED)
        else:
            _run(paragraph, part, size=size)


# ----------------------------------------------------------------------------
# Builder utama
# ----------------------------------------------------------------------------
def build_pptx(d: FormData) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(11.69)   # A4 landscape
    prs.slide_height = Inches(8.27)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout kosong

    # ---------------- Header ----------------
    if d.logo_kiri:
        raw, w, h = _prepare_image(d.logo_kiri, 500, keep_alpha=True)
        x, y, ww, hh = _fit(w, h, 0.4, 0.25, 2.2, 0.7, "left")
        slide.shapes.add_picture(io.BytesIO(raw), Inches(x), Inches(y), Inches(ww), Inches(hh))
    else:
        _, tf = _textbox(slide, 0.4, 0.2, 2.2, 0.8)
        _run(tf.paragraphs[0], "BUMN", size=20, bold=True, color=NAVY)
        _run(tf.add_paragraph(), "UNTUK INDONESIA", size=8, bold=True, color=NAVY)

    if d.logo_kanan:
        raw, w, h = _prepare_image(d.logo_kanan, 500, keep_alpha=True)
        x, y, ww, hh = _fit(w, h, 9.29, 0.25, 2.0, 0.7, "right")
        slide.shapes.add_picture(io.BytesIO(raw), Inches(x), Inches(y), Inches(ww), Inches(hh))
    else:
        _simple_text(slide, "PERTAMINA", 9.0, 0.25, 2.3, 0.7, size=16, bold=True,
                     color=NAVY, align=PP_ALIGN.RIGHT)

    _simple_text(slide, "FORM PENJELASAN DOKUMENTASI", 2.4, 0.4, 6.9, 0.5,
                 size=18, bold=True, align=PP_ALIGN.CENTER)

    # ---------------- Jenis Kegiatan ----------------
    _rect(slide, 0.4, 1.02, 10.89, 0.30, SOFT)
    _row(slide, "Jenis Kegiatan", d.jenis_kegiatan, 0.45, 1.04, 1.25, 9.5)

    # ---------------- Data Narasumber ----------------
    _rect(slide, 0.4, 1.45, 6.6, 0.26, NAVY)
    _simple_text(slide, "Data Narasumber", 0.45, 1.45, 6.4, 0.26, size=11, bold=True, color=WHITE)
    _rect(slide, 0.4, 1.71, 6.6, 1.90, SOFT)
    rows = [
        ("Nama Lengkap", d.nama, False),
        ("NIK", d.nik, False),
        ("Jabatan", d.jabatan, False),
        ("Email", d.email, True),
        ("HSH", d.hsh, False),
        ("Perusahaan", d.perusahaan, False),
        ("Direktorat", d.direktorat, False),
    ]
    for i, (lab, val, is_mail) in enumerate(rows):
        _row(slide, lab, val, 0.45, 1.74 + i * 0.263, 1.25, 5.25, email=is_mail)

    # ---------------- Jadwal Pelaksanaan ----------------
    _rect(slide, 7.1, 1.45, 4.19, 0.26, NAVY)
    _simple_text(slide, "Jadwal Pelaksanaan", 7.15, 1.45, 4.0, 0.26, size=11, bold=True, color=WHITE)
    _rect(slide, 7.1, 1.71, 4.19, 0.62, SOFT)
    _row(slide, "Hari/Tanggal", d.tanggal, 7.15, 1.74, 1.25, 2.85)
    waktu = f"{d.waktu_mulai.replace(':', '.')} WIB  s/d  {d.waktu_selesai.replace(':', '.')} WIB"
    _row(slide, "Waktu (start - end)", waktu, 7.15, 2.00, 1.25, 2.85)

    # ---------------- Informasi ----------------
    _rect(slide, 0.4, 3.75, 10.89, 0.26, NAVY)
    _simple_text(slide, "Informasi", 0.45, 3.75, 5.0, 0.26, size=11, bold=True, color=WHITE)
    _rect(slide, 0.4, 4.01, 10.89, 3.98, SOFT)
    _simple_text(slide, f"Topik/ Judul Materi : {d.topik}", 0.45, 4.03, 10.7, 0.28, size=9.5)

    # Uraian singkat (kiri)
    words = count_words(strip_markup(d.uraian))
    size = 8.5 if words > 200 else 9 if words > 150 else 10
    _, tf = _textbox(slide, 0.45, 4.34, 4.95, 3.60, anchor=MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.space_after = Pt(3)
    _run(p, "Uraian Singkat :", size=11)
    for para in [s.strip() for s in re.split(r"\n+", d.uraian) if s.strip()]:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        _add_markup_runs(p, para, size)

    # Gambar + deskripsi (kanan)
    imgs = d.images
    if imgs:
        area_x, area_y, area_w, area_h = 5.6, 4.36, 5.6, 3.58
        slot = area_h / len(imgs)
        gap = 0.06
        for i, im in enumerate(imgs):
            y0 = area_y + i * slot
            top = y0
            if im.caption.strip():
                r = _rect(slide, area_x, y0, area_w, 0.30, YELLOW)
                tf = r.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = tf.margin_right = Inches(0.05)
                tf.margin_top = tf.margin_bottom = Inches(0.02)
                pp = tf.paragraphs[0]
                pp.alignment = PP_ALIGN.CENTER
                _run(pp, im.caption.strip(), size=8, bold=True, color=RGBColor(0, 0, 0))
                top = y0 + 0.34
            box_h = slot - (top - y0) - gap
            _rect(slide, area_x, top, area_w, box_h, WHITE)
            raw, w, h = _prepare_image(im.data, 1600)
            x, y, ww, hh = _fit(w, h, area_x, top, area_w, box_h)
            slide.shapes.add_picture(io.BytesIO(raw), Inches(x), Inches(y), Inches(ww), Inches(hh))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
